"""Rilevamento di password e credenziali nei documenti."""
import json

import pytest
from conftest import csrf, login

from app import secretscan
from app.config import settings
from app.secretscan import mask, redact, scan_line, scan_text

POSITIVE = [
    "Password: Estate2026!",
    "La password del wifi è Ciao1234",
    "pwd = s3cr3t#",
    "**Password:** `Xy7!kL9q`",
    "PIN: 48213",
    "Chiave WPA2: CasaRossa_2024",
    "API key: sk-proj-abcDEF1234567890xyz",
    "Credenziali: admin / Pa55word!",
    "Password del database di produzione: Pr0d!Db",
    'password="SuperSegreta99"',
    "Il codice di accesso è 774412",
    "postgres://app:Xyz789!@db.interno:5432/app",
    "token: ghp_abcdefghijklmnopqrstuvwxyz0123456789",
    "AKIAIOSFODNN7EXAMPLE",
    "-----BEGIN RSA PRIVATE KEY-----",
    "parola d'ordine: Girasole77",
]
NEGATIVE = [
    "Chiedi la password all'IT",
    "La password è obbligatoria",
    "Password: richiedila al supporto",
    "password: ********",
    "Password: <la tua password>",
    "password: ${DB_PASSWORD}",
    "La password deve avere almeno 8 caratteri.",
    "Reimposta la password dal portale https://sso.example.com",
    "Token: da generare in Impostazioni",
    "PIN: 2024",
    "Il PIN è personale",
    "Password dimenticata? Clicca qui.",
    "Password: Obbligatoria",
    "password: (non condividerla)",
    "https://user:${PASSWORD}@host/x",
    "Accesso: tramite SSO",
]


@pytest.mark.parametrize("line", POSITIVE)
def test_trova_password(line):
    assert scan_line(line), line


@pytest.mark.parametrize("line", NEGATIVE)
def test_ignora_testo_normale(line):
    assert not scan_line(line), [(f.label, f.masked) for f in scan_line(line)]


def test_maschera_e_contesto_senza_valore():
    f = scan_line("Password VPN: Estate2026!")[0]
    assert f.masked == "E•••••••••!" and "Estate2026!" not in f.context and f.masked in f.context
    assert mask("abc") == "••••"


def test_tabella_markdown():
    md = "| Servizio | Utente | Password |\n|---|---|---|\n| VPN | mario | estate2024 |\n| Wiki | anna | ******** |\n"
    fs = scan_text(md)
    assert [(f.line, f.label) for f in fs] == [(3, "Password in una tabella")]


def test_redact():
    out = redact("Accesso VPN. Password: Estate2026! Server: postgres://app:Xyz789!@db:5432")
    assert "Estate2026!" not in out and "Xyz789!" not in out and "••••••" in out


def test_foglio_excel_con_colonna_password(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Accessi"
    ws.append(["Servizio", "Utente", "Password"])
    ws.append(["VPN", "mario", "estate2024"])
    ws.append(["Posta", "anna", "-"])
    ws2 = wb.create_sheet("Note")
    ws2.append(["PIN", "4821"])
    p = tmp_path / "accessi.xlsx"
    wb.save(p)
    fs = secretscan.scan_file(p)
    locs = sorted(f.location for f in fs)
    assert locs == ["Foglio Accessi, riga 2", "Foglio Note, riga 1"]
    assert any(f.anchor == "?sheet=Accessi" for f in fs)


def test_tabella_word(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("Accessi ai sistemi")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Utente", "Password"
    t.cell(1, 0).text, t.cell(1, 1).text = "mario", "Sole2026!"
    p = tmp_path / "accessi.docx"
    d.save(p)
    fs = secretscan.scan_file(p)
    assert fs and fs[0].location == "Tabella 1, riga 2"


# ------------------------------------------------------------------ integrazione
def _save(client, token, content, confirm=""):
    return client.post("/edit/01-Guida/01-primi-passi.md",
                       data={"csrf": token, "content": content, "confirm_secrets": confirm}, follow_redirects=False)


def test_editor_chiede_conferma(client, data_dir):
    login(client, "editor", "ailadoc-editor")
    token = csrf(client)
    page = data_dir / "docs/01-Guida/01-primi-passi.md"
    before = page.read_text()
    r = _save(client, token, "# Primi passi\n\nWiFi ospiti - password: Ciao1234!\n")
    assert r.status_code == 422 and "possibile password" in r.text and "Ciao1234!" in r.text  # il testo resta nell'editor
    assert page.read_text() == before
    r = _save(client, token, "# Primi passi\n\nWiFi ospiti - password: Ciao1234!\n", confirm="1")
    assert r.status_code == 303 and "Ciao1234!" in page.read_text()
    # la pagina mostra l'avviso all'editor (con il valore mascherato)...
    html = client.get("/d/01-Guida/01-primi-passi.md").text
    assert "secret-banner" in html and "C•••••••!" in html and "Riga 3" in html
    # ...ma non a chi può solo leggere
    login(client, "viewer", "ailadoc-viewer")
    assert "secret-banner" not in client.get("/d/01-Guida/01-primi-passi.md").text


def test_policy_block(client, data_dir, monkeypatch):
    monkeypatch.setattr(settings, "secrets_policy", "block")
    login(client, "editor", "ailadoc-editor")
    token = csrf(client)
    r = _save(client, token, "Password: Ciao1234!", confirm="1")
    assert r.status_code == 422 and "toglile prima di salvare" in r.text
    assert "Ciao1234!" not in (data_dir / "docs/01-Guida/01-primi-passi.md").read_text()
    r = client.post("/upload", data={"csrf": token, "folder": "01-Guida"}, headers={"accept": "application/json"},
                    files=[("files", ("note.txt", b"utente: mario\npassword: Ciao1234!\n", "text/plain"))])
    assert r.json()["saved"] == [] and "password" in r.json()["errors"][0]
    assert not (data_dir / "docs/01-Guida/note.txt").exists()


def test_upload_avvisa(client, data_dir):
    login(client, "editor", "ailadoc-editor")
    token = csrf(client)
    r = client.post("/upload", data={"csrf": token, "folder": "01-Guida"}, headers={"accept": "application/json"},
                    files=[("files", ("note.txt", b"utente: mario\npassword: Ciao1234!\n", "text/plain"))])
    d = r.json()
    assert d["saved"] == ["01-Guida/note.txt"] and d["secrets"] == [{"path": "01-Guida/note.txt", "count": 1}]


def test_check_live_e_falso_positivo(client, data_dir):
    login(client, "editor", "ailadoc-editor")
    token = csrf(client)
    h = {"x-csrf-token": token}
    content = "Codice progetto: PRJ2026x\nPassword: Ciao1234!"
    d = client.post("/api/secrets/check", json={"content": content, "path": "01-Guida/01-primi-passi.md"}, headers=h).json()
    assert [f["line"] for f in d["findings"]] == [2]
    fp = d["findings"][0]["fp"]
    assert "Ciao1234!" not in json.dumps(d)
    client.post("/api/secrets/ignore", json={"path": "01-Guida/01-primi-passi.md", "fp": fp}, headers=h)
    d = client.post("/api/secrets/check", json={"content": content, "path": "01-Guida/01-primi-passi.md"}, headers=h).json()
    assert d["findings"] == []
    # un editor non può ignorare un valore "ovunque"
    assert client.post("/api/secrets/ignore", json={"path": "*", "fp": fp}, headers=h).status_code == 403


def test_pagina_admin_e_file_copiati_a_mano(client, data_dir):
    (data_dir / "docs/01-Guida/accessi.md").write_text("# Accessi\n\nVPN: utente mario, password Estate2026!\n\nPassword: Sole2026!\n")
    (data_dir / "docs/02-Procedure/server.txt").write_text("db = postgres://app:Xyz789!@db:5432/app\n")
    login(client, "admin", "ailadoc-admin")
    token = csrf(client)
    html = client.get("/admin/secrets").text
    assert "/01-Guida/accessi.md" in html and "/02-Procedure/server.txt" in html
    assert "Estate2026!" not in html and "Sole2026!" not in html and "Xyz789!" not in html
    assert "nav-secret" in client.get("/").text  # chiave accanto ai documenti nel menu dell'admin
    # la cache su disco non contiene i valori
    cache = (data_dir / "cache/secrets.json").read_text()
    assert "Sole2026!" not in cache and "Xyz789!" not in cache
    # falso positivo "ovunque" -> sparisce
    fp = secretscan.SCANNER.findings("02-Procedure/server.txt")[0].fp
    client.post("/api/secrets/ignore", json={"path": "*", "fp": fp, "masked": "X•••", "label": "x"},
                headers={"x-csrf-token": token})
    html = client.get("/admin/secrets").text
    assert "/02-Procedure/server.txt</strong>" not in html and "tutti i documenti" in html


def test_ricerca_e_assistente_non_mostrano_password(client, data_dir):
    (data_dir / "docs/01-Guida/wifi.md").write_text("# WiFi ospiti\n\nLa rete si chiama Ospiti. Password: Ciao1234! valida fino a dicembre.\n")
    login(client, "viewer", "ailadoc-viewer")
    assert "Ciao1234!" not in client.get("/search?q=ospiti").text
    from app import assistant

    assistant.INDEX.refresh(force=True)
    r = client.post("/api/assistant", json={"question": "password wifi ospiti"}, headers={"x-csrf-token": csrf(client)})
    body = json.dumps(r.json(), ensure_ascii=False)
    assert "wifi.md" in body and "Ciao1234!" not in body
