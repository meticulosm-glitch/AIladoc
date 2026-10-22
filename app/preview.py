"""Anteprime: conversione Office->PDF via LibreOffice e fogli di calcolo come tabelle HTML."""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path

from . import metrics
from .config import settings

log = logging.getLogger("ailadoc.preview")
_lock = threading.Lock()
OFFICE_PDF_EXT = {"ppt", "pptx", "pps", "ppsx", "odp", "doc", "docx", "odt", "rtf", "xls", "ods"}
MAX_ROWS = 2000
MAX_COLS = 60


def soffice_bin() -> str | None:
    if not settings.office_preview:
        return None
    return shutil.which("soffice") or shutil.which("libreoffice")


def can_pdf_preview(name: str) -> bool:
    return name.rsplit(".", 1)[-1].lower() in OFFICE_PDF_EXT and soffice_bin() is not None


def _cache_file(src: Path) -> Path:
    st = src.stat()
    key = hashlib.sha256(f"{src}|{st.st_mtime_ns}|{st.st_size}".encode()).hexdigest()[:32]
    return settings.cache_dir / "pdf" / f"{key}.pdf"


def office_to_pdf(src: Path) -> Path | None:
    """Converte (con cache) un file Office in PDF. Bloccante: chiamare in un thread."""
    binary = soffice_bin()
    if not binary:
        return None
    out = _cache_file(src)
    if out.exists():
        metrics.OFFICE_CACHE.labels(result="hit").inc()
        return out
    out.parent.mkdir(parents=True, exist_ok=True)
    with _lock:  # LibreOffice non gradisce istanze concorrenti sullo stesso profilo
        if out.exists():
            metrics.OFFICE_CACHE.labels(result="hit").inc()
            return out
        metrics.OFFICE_CACHE.labels(result="miss").inc()
        with tempfile.TemporaryDirectory() as tmp, \
                metrics.timed(metrics.OFFICE_CONVERSION, outcome="ok") as t:
            tmpd = Path(tmp)
            local = tmpd / ("input." + src.name.rsplit(".", 1)[-1].lower())
            shutil.copyfile(src, local)
            cmd = [binary, f"-env:UserInstallation=file://{tmpd}/profile", "--headless", "--norestore",
                   "--convert-to", "pdf", "--outdir", str(tmpd), str(local)]
            try:
                subprocess.run(cmd, check=True, timeout=settings.office_timeout,
                               stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
            except (subprocess.SubprocessError, OSError) as e:
                t["outcome"] = "timeout" if isinstance(e, subprocess.TimeoutExpired) else "error"
                log.warning("Conversione fallita per %s: %s", src, e)
                return None
            produced = tmpd / "input.pdf"
            if not produced.exists():
                t["outcome"] = "error"
                return None
            shutil.move(str(produced), out)
    return out


def sheet_tables(src: Path) -> list[dict]:
    """Ritorna [{name, rows: [[...]], truncated}] per xlsx/xlsm/csv/tsv."""
    ext = src.suffix.lower().lstrip(".")
    if ext in ("csv", "tsv"):
        raw = src.read_bytes()
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        text = text or ""
        if ext == "tsv":
            delim = "\t"
        else:
            try:
                delim = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
            except csv.Error:
                delim = ","
        rows, truncated = [], False
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
            if i >= MAX_ROWS:
                truncated = True
                break
            rows.append(row[:MAX_COLS])
        return [{"name": src.name, "rows": rows, "truncated": truncated}]
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(src, read_only=True, data_only=True)
        out = []
        try:
            for ws in wb.worksheets:
                rows, truncated = [], False
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= MAX_ROWS:
                        truncated = True
                        break
                    rows.append(["" if v is None else _fmt(v) for v in row[:MAX_COLS]])
                while rows and not any(c != "" for c in rows[-1]):
                    rows.pop()
                width = max((max((j + 1 for j, c in enumerate(r) if c != ""), default=0) for r in rows), default=0)
                rows = [r[:width] for r in rows]
                out.append({"name": ws.title, "rows": rows, "truncated": truncated})
        finally:
            wb.close()
        return out
    return []


def _fmt(v) -> str:
    import datetime as dt

    if isinstance(v, float):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if v != int(v) else f"{int(v)}"
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y %H:%M") if (v.hour or v.minute) else v.strftime("%d/%m/%Y")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    return str(v)
