"""Motore delle tabelle dinamiche: caricamento, espressioni sicure, filtri, raggruppamenti, pivot.

Le espressioni (``filter``, ``compute``, ``highlight``) sono valutate da un interprete minimo
basato su ``ast``: niente ``eval``, solo operatori aritmetici/confronto/logici, le colonne
della tabella e un elenco chiuso di funzioni.
"""
from __future__ import annotations

import ast
import csv
import datetime as dt
import io
import math
import re
import statistics
import threading
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

MAX_ROWS = 5000
MAX_COLS = 60
MAX_OUT_ROWS = 2000


class TableError(ValueError):
    pass


@dataclass
class Table:
    columns: list[str]
    rows: list[list[Any]] = field(default_factory=list)
    truncated: bool = False

    def index(self, name: str) -> int:
        key = _norm(name)
        for i, c in enumerate(self.columns):
            if _norm(c) == key:
                return i
        raise TableError(f"Colonna «{name}» non trovata. Colonne disponibili: {', '.join(self.columns)}")

    def has(self, name: str) -> bool:
        key = _norm(name)
        return any(_norm(c) == key for c in self.columns)

    def col(self, name: str) -> list[Any]:
        i = self.index(name)
        return [r[i] if i < len(r) else None for r in self.rows]

    def is_numeric(self, i: int) -> bool:
        vals = [r[i] for r in self.rows if i < len(r) and r[i] not in (None, "")]
        return bool(vals) and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals)


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s)).strip().lower()


# --------------------------------------------------------------------------- valori
_NUM_IT = re.compile(r"^[-+]?\d{1,3}(?:\.\d{3})+(?:,\d+)?$|^[-+]?\d+,\d+$")
_NUM_EN = re.compile(r"^[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?$|^[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?$")
_DATE_IT = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}):(\d{2}))?$")
_DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[T\s](\d{2}):(\d{2}))?")


def parse_value(v: Any) -> Any:
    """Testo di una cella -> numero/data se riconoscibile (formati italiani e inglesi), altrimenti testo."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return int(v) if isinstance(v, float) and v.is_integer() and abs(v) < 1e15 else v
    if isinstance(v, dt.datetime):
        return v.date() if not (v.hour or v.minute or v.second) else v
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    if not s:
        return None
    t = s.replace(" ", "").replace(" ", "")
    pct = t.endswith("%")
    t = t.rstrip("%").replace("€", "").replace("$", "").replace("EUR", "")
    num = None
    if _NUM_IT.match(t):
        num = float(t.replace(".", "").replace(",", "."))
    elif _NUM_EN.match(t):
        num = float(t.replace(",", ""))
    if num is not None:
        if pct:
            num = num / 100
        return int(num) if num.is_integer() and not pct and abs(num) < 1e15 else num
    m = _DATE_IT.match(s)
    if m:
        try:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if m.group(4):
                return dt.datetime(y, mo, d, int(m.group(4)), int(m.group(5)))
            return dt.date(y, mo, d)
        except ValueError:
            return s
    m = _DATE_ISO.match(s)
    if m and len(s) <= 19:
        try:
            y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if m.group(4):
                return dt.datetime(y, mo, d, int(m.group(4)), int(m.group(5)))
            return dt.date(y, mo, d)
        except ValueError:
            return s
    return s


def _from_rows(rows: list[list[Any]], header: int = 1, truncated: bool = False) -> Table:
    rows = [list(r) for r in rows]
    if len(rows) < header:
        raise TableError("La tabella è vuota.")
    head = rows[header - 1]
    body = rows[header:]
    width = max([len(head)] + [len(r) for r in body[:200]]) if (head or body) else 0
    width = min(width, MAX_COLS)
    cols, seen = [], {}
    for i in range(width):
        name = str(head[i]).strip() if i < len(head) and head[i] not in (None, "") else f"Colonna {i + 1}"
        if name.lower() in seen:
            seen[name.lower()] += 1
            name = f"{name} ({seen[name.lower()]})"
        else:
            seen[name.lower()] = 1
        cols.append(name)
    out = []
    for r in body:
        vals = [parse_value(r[i]) if i < len(r) else None for i in range(width)]
        if any(v not in (None, "") for v in vals):
            out.append(vals)
    return Table(cols, out, truncated)


# --------------------------------------------------------------------------- caricamento
_cache: OrderedDict = OrderedDict()
_cache_lock = threading.Lock()


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_csv_text(text: str, delim: str | None = None) -> list[list[str]]:
    if delim is None:
        sample = text[:5000]
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delim = "\t" if "\t" in sample else (";" if sample.count(";") > sample.count(",") else ",")
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
        if i > MAX_ROWS:
            break
        rows.append(row[:MAX_COLS])
    return rows


def parse_markdown_table(text: str) -> list[list[str]] | None:
    lines = [ln.strip() for ln in text.strip().splitlines() if ln.strip()]
    if len(lines) < 2 or not all(ln.startswith("|") or ln.count("|") >= 1 for ln in lines):
        return None
    if not re.match(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$", lines[1]):
        return None
    rows = []
    for i, ln in enumerate(lines):
        if i == 1:
            continue
        ln = ln.strip()
        if ln.startswith("|"):
            ln = ln[1:]
        if ln.endswith("|") and not ln.endswith("\\|"):
            ln = ln[:-1]
        cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", ln)]
        rows.append(cells)
    return rows


def load_inline(text: str) -> Table:
    text = text.strip("\n")
    if not text.strip():
        raise TableError("Indica un file con «source:» oppure scrivi i dati dentro il blocco (CSV o tabella Markdown).")
    rows = parse_markdown_table(text)
    if rows is None:
        rows = parse_csv_text(text)
    return _from_rows(rows)


def load_file(p: Path, sheet: str = "", header: int = 1) -> tuple[Table, str]:
    """Ritorna (tabella, nome del foglio usato). Cache in memoria per (file, mtime, foglio)."""
    st = p.stat()
    key = (str(p), st.st_mtime_ns, st.st_size, sheet, header)
    with _cache_lock:
        if key in _cache:
            _cache.move_to_end(key)
            t, name = _cache[key]
            return Table(list(t.columns), [list(r) for r in t.rows], t.truncated), name
    ext = p.suffix.lower().lstrip(".")
    truncated = False
    if ext in ("csv", "tsv", "txt"):
        rows = parse_csv_text(_decode(p.read_bytes()), "\t" if ext == "tsv" else None)
        truncated = len(rows) > MAX_ROWS
        name = p.name
    elif ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            ws = _pick_sheet(wb.worksheets, sheet)
            name = ws.title
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS:
                    truncated = True
                    break
                rows.append(list(row[:MAX_COLS]))
        finally:
            wb.close()
        while rows and not any(c not in (None, "") for c in rows[-1]):
            rows.pop()
    else:
        raise TableError("Formato non supportato: usa .xlsx, .xlsm, .csv o .tsv.")
    t = _from_rows(rows[: MAX_ROWS + header], header, truncated)
    with _cache_lock:
        _cache[key] = (t, name)
        while len(_cache) > 32:
            _cache.popitem(last=False)
    return Table(list(t.columns), [list(r) for r in t.rows], t.truncated), name


def _pick_sheet(sheets, sheet: str):
    if not sheet:
        return sheets[0]
    s = str(sheet).strip()
    for ws in sheets:
        if ws.title.lower() == s.lower():
            return ws
    if s.isdigit() and 1 <= int(s) <= len(sheets):
        return sheets[int(s) - 1]
    raise TableError(f"Foglio «{sheet}» non trovato. Fogli: {', '.join(ws.title for ws in sheets)}")


# --------------------------------------------------------------------------- espressioni
def _num(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    if v is None or v == "":
        return None
    p = parse_value(v)
    return p if isinstance(p, (int, float)) else None


def _days(a, b=None):
    a = a if isinstance(a, (dt.date, dt.datetime)) else parse_value(a)
    b = dt.date.today() if b is None else (b if isinstance(b, (dt.date, dt.datetime)) else parse_value(b))
    if not isinstance(a, dt.date) or not isinstance(b, dt.date):
        return None
    a = a.date() if isinstance(a, dt.datetime) else a
    b = b.date() if isinstance(b, dt.datetime) else b
    return (a - b).days


def _date_part(attr):
    def f(v):
        v = v if isinstance(v, (dt.date, dt.datetime)) else parse_value(v)
        return getattr(v, attr) if isinstance(v, dt.date) else None
    return f


def _quarter(v):
    m = _date_part("month")(v)
    return f"T{(m - 1) // 3 + 1}" if m else None


def _round(v, n=0):
    v = _num(v)
    return None if v is None else (round(v, int(n)) if n else round(v))


def _text(v):
    return "" if v is None else str(v)


FUNCS = {
    "round": _round, "arrotonda": _round,
    "abs": lambda v: None if _num(v) is None else abs(_num(v)),
    "min": lambda *a: min((x for x in a if x is not None), default=None),
    "max": lambda *a: max((x for x in a if x is not None), default=None),
    "len": lambda v: len(_text(v)), "lower": lambda v: _text(v).lower(), "upper": lambda v: _text(v).upper(),
    "contains": lambda a, b: _text(b).lower() in _text(a).lower(),
    "contiene": lambda a, b: _text(b).lower() in _text(a).lower(),
    "startswith": lambda a, b: _text(a).lower().startswith(_text(b).lower()),
    "inizia": lambda a, b: _text(a).lower().startswith(_text(b).lower()),
    "year": _date_part("year"), "anno": _date_part("year"), "month": _date_part("month"),
    "mese": _date_part("month"), "day": _date_part("day"), "giorno": _date_part("day"),
    "quarter": _quarter, "trimestre": _quarter,
    "today": lambda: dt.date.today(), "oggi": lambda: dt.date.today(),
    "days": _days, "giorni": _days,
    "if": lambda c, a, b=None: a if c else b, "se": lambda c, a, b=None: a if c else b,
    "concat": lambda *a: "".join(_text(x) for x in a), "unisci": lambda *a: "".join(_text(x) for x in a),
    "coalesce": lambda *a: next((x for x in a if x not in (None, "")), None),
    "num": _num, "text": _text, "testo": _text,
    "sqrt": lambda v: None if _num(v) is None or _num(v) < 0 else math.sqrt(_num(v)),
}

_STR_SPLIT = re.compile(r"(\"(?:[^\"\\]|\\.)*\"|'(?:[^'\\]|\\.)*')")
_BRACKET = re.compile(r"\[([^\[\]]+)\]")


class Expr:
    """Espressione compilata, valutabile su una riga della tabella."""

    def __init__(self, src: str, table: Table):
        self.src = str(src).strip()
        if len(self.src) > 1000:
            raise TableError("Espressione troppo lunga.")
        self.cols: dict[str, int] = {}
        code = self._prepare(self.src, table)
        try:
            self.tree = ast.parse(code, mode="eval")
        except SyntaxError as e:
            raise TableError(f"Espressione non valida: {self.src}") from e
        self._check(self.tree, table)

    def _prepare(self, src: str, table: Table) -> str:
        parts = _STR_SPLIT.split(src)
        out = []
        for i, part in enumerate(parts):
            if i % 2 == 1:  # stringa letterale
                out.append(part)
                continue

            def br(m):
                idx = table.index(m.group(1))
                name = f"__c{idx}"
                self.cols[name] = idx
                return name

            part = _BRACKET.sub(br, part)
            part = re.sub(r"\bAND\b", " and ", part)
            part = re.sub(r"\bOR\b", " or ", part)
            part = re.sub(r"\bNOT\b", " not ", part)
            part = part.replace("<>", "!=")
            part = re.sub(r"(?<![<>=!])=(?!=)", "==", part)
            out.append(part)
        return "".join(out)

    _OK = (ast.Expression, ast.BoolOp, ast.BinOp, ast.UnaryOp, ast.Compare, ast.IfExp, ast.Call, ast.Name,
           ast.Constant, ast.Load, ast.And, ast.Or, ast.Not, ast.USub, ast.UAdd, ast.Add, ast.Sub, ast.Mult,
           ast.Div, ast.FloorDiv, ast.Mod, ast.Pow, ast.Eq, ast.NotEq, ast.Lt, ast.LtE, ast.Gt, ast.GtE, ast.In,
           ast.NotIn, ast.Tuple, ast.List)

    def _check(self, tree, table: Table) -> None:
        for node in ast.walk(tree):
            if not isinstance(node, self._OK):
                raise TableError(f"Costrutto non ammesso nell'espressione: {self.src}")
            if isinstance(node, ast.Call):
                if not isinstance(node.func, ast.Name) or node.func.id.lower() not in FUNCS or node.keywords:
                    fn = getattr(node.func, "id", "?")
                    raise TableError(f"Funzione «{fn}» non disponibile. Funzioni: {', '.join(sorted(FUNCS))}")
            elif isinstance(node, ast.Name) and node.id not in self.cols:
                if node.id.lower() in FUNCS or node.id in ("True", "False", "None", "true", "false", "vero", "falso"):
                    continue
                idx = table.index(node.id)  # nome di colonna senza spazi
                self.cols[node.id] = idx

    def __call__(self, row: list[Any]) -> Any:
        try:
            return self._ev(self.tree.body, row)
        except (TypeError, ValueError, ZeroDivisionError, OverflowError, ArithmeticError):
            return None

    def _ev(self, n, row):  # noqa: C901 - interprete
        if isinstance(n, ast.Constant):
            return n.value
        if isinstance(n, ast.Name):
            if n.id in self.cols:
                i = self.cols[n.id]
                return row[i] if i < len(row) else None
            low = n.id.lower()
            return {"true": True, "false": False, "vero": True, "falso": False, "none": None}.get(low)
        if isinstance(n, (ast.Tuple, ast.List)):
            return [self._ev(e, row) for e in n.elts]
        if isinstance(n, ast.BoolOp):
            if isinstance(n.op, ast.And):
                return all(self._ev(v, row) for v in n.values)
            return any(self._ev(v, row) for v in n.values)
        if isinstance(n, ast.UnaryOp):
            v = self._ev(n.operand, row)
            if isinstance(n.op, ast.Not):
                return not v
            v = _num(v)
            return None if v is None else (-v if isinstance(n.op, ast.USub) else v)
        if isinstance(n, ast.IfExp):
            return self._ev(n.body, row) if self._ev(n.test, row) else self._ev(n.orelse, row)
        if isinstance(n, ast.Call):
            return FUNCS[n.func.id.lower()](*[self._ev(a, row) for a in n.args])
        if isinstance(n, ast.BinOp):
            a, b = self._ev(n.left, row), self._ev(n.right, row)
            if isinstance(n.op, ast.Add) and (isinstance(a, str) or isinstance(b, str)):
                return _text(a) + _text(b)
            if isinstance(n.op, ast.Sub) and isinstance(a, dt.date) and isinstance(b, dt.date):
                return _days(a, b)
            a, b = _num(a), _num(b)
            if a is None or b is None:
                return None
            if isinstance(n.op, ast.Pow) and abs(b) > 100:
                raise ValueError("esponente troppo grande")
            if isinstance(n.op, ast.Add):
                return a + b
            if isinstance(n.op, ast.Sub):
                return a - b
            if isinstance(n.op, ast.Mult):
                return a * b
            if isinstance(n.op, ast.Div):
                return a / b
            if isinstance(n.op, ast.FloorDiv):
                return a // b
            if isinstance(n.op, ast.Mod):
                return a % b
            if isinstance(n.op, ast.Pow):
                return a ** b
        if isinstance(n, ast.Compare):
            left = self._ev(n.left, row)
            for op, comp in zip(n.ops, n.comparators):
                right = self._ev(comp, row)
                if not _compare(op, left, right):
                    return False
                left = right
            return True
        raise ValueError("nodo non gestito")


def _cmp_pair(a, b):
    """Rende confrontabili numeri scritti come testo, date e testo (senza maiuscole)."""
    if isinstance(a, (int, float)) and isinstance(b, str):
        b = _num(b) if _num(b) is not None else b
    if isinstance(b, (int, float)) and isinstance(a, str):
        a = _num(a) if _num(a) is not None else a
    if isinstance(a, dt.date) and isinstance(b, str):
        b = parse_value(b)
    if isinstance(b, dt.date) and isinstance(a, str):
        a = parse_value(a)
    if isinstance(a, dt.datetime) and not isinstance(b, dt.datetime) and isinstance(b, dt.date):
        a = a.date()
    if isinstance(b, dt.datetime) and not isinstance(a, dt.datetime) and isinstance(a, dt.date):
        b = b.date()
    if isinstance(a, str) and isinstance(b, str):
        return a.lower(), b.lower()
    return a, b


def _compare(op, a, b) -> bool:
    if isinstance(op, (ast.In, ast.NotIn)):
        if isinstance(b, list):
            hit = any(_cmp_pair(a, x)[0] == _cmp_pair(a, x)[1] for x in b)
        else:
            hit = _text(a).lower() in _text(b).lower()
        return hit if isinstance(op, ast.In) else not hit
    a, b = _cmp_pair(a, b)
    if isinstance(op, ast.Eq):
        return a == b
    if isinstance(op, ast.NotEq):
        return a != b
    if a is None or b is None:
        return False
    try:
        if isinstance(op, ast.Lt):
            return a < b
        if isinstance(op, ast.LtE):
            return a <= b
        if isinstance(op, ast.Gt):
            return a > b
        if isinstance(op, ast.GtE):
            return a >= b
    except TypeError:
        return False
    return False


# --------------------------------------------------------------------------- operazioni
AGG_ALIASES = {
    "sum": "sum", "somma": "sum", "totale": "sum", "avg": "avg", "mean": "avg", "media": "avg",
    "min": "min", "minimo": "min", "max": "max", "massimo": "max", "count": "count", "conta": "count",
    "conteggio": "count", "count_distinct": "count_distinct", "distinct": "count_distinct", "distinti": "count_distinct",
    "median": "median", "mediana": "median", "first": "first", "primo": "first", "last": "last", "ultimo": "last",
}
AGG_LABEL = {"sum": "somma", "avg": "media", "min": "min", "max": "max", "count": "conteggio",
             "count_distinct": "distinti", "median": "mediana", "first": "primo", "last": "ultimo"}


def aggregate(func: str, values: list[Any]) -> Any:
    f = AGG_ALIASES.get(str(func).strip().lower())
    if f is None:
        raise TableError(f"Funzione di aggregazione «{func}» sconosciuta. Usa: {', '.join(sorted(set(AGG_ALIASES.values())))}")
    if f == "count":
        return sum(1 for v in values if v not in (None, ""))
    if f == "count_distinct":
        return len({v for v in values if v not in (None, "")})
    if f == "first":
        return next((v for v in values if v not in (None, "")), None)
    if f == "last":
        return next((v for v in reversed(values) if v not in (None, "")), None)
    nums = [n for n in (_num(v) for v in values) if n is not None]
    if f in ("min", "max") and not nums:
        vals = [v for v in values if v not in (None, "")]
        try:
            return (min if f == "min" else max)(vals) if vals else None
        except TypeError:
            return None
    if not nums:
        return None
    if f == "sum":
        s = math.fsum(nums)
        return int(s) if float(s).is_integer() and all(isinstance(n, int) for n in nums) else s
    if f == "avg":
        return math.fsum(nums) / len(nums)
    if f == "min":
        return min(nums)
    if f == "max":
        return max(nums)
    if f == "median":
        return statistics.median(nums)
    return None


_AGG_SPEC = re.compile(r"^\s*(\w+)\s*\(\s*(.*?)\s*\)\s*(?:(?:as|come)\s+(.+?))?\s*$", re.I)


def parse_agg_specs(spec: Any, table: Table) -> list[tuple[str, str, str | None]]:
    """-> [(funzione, colonna_sorgente|None, nome_output)]"""
    out = []
    if isinstance(spec, dict):
        for k, v in spec.items():
            k, v = str(k), str(v)
            m = _AGG_SPEC.match(v)
            if m:  # Alias: sum(Colonna)
                out.append((m.group(1), m.group(2) or None, k))
            elif table.has(k):  # Colonna: funzione
                out.append((v, k, k))
            elif AGG_ALIASES.get(v.lower()) == "count":  # Ordini: count
                out.append((v, None, k))
            else:
                raise TableError(f"Colonna «{k}» non trovata per l'aggregazione.")
    elif isinstance(spec, (list, tuple, str)):
        items = [spec] if isinstance(spec, str) else spec
        for it in items:
            for piece in str(it).split(";") if isinstance(it, str) and ";" in str(it) else [it]:
                m = _AGG_SPEC.match(str(piece))
                if not m:
                    raise TableError(f"Aggregazione non valida: «{piece}». Esempio: sum(Totale) as Fatturato")
                func, col, alias = m.group(1), m.group(2) or None, m.group(3)
                if col in ("*", ""):
                    col = None
                out.append((func, col, alias or (f"{col} ({AGG_LABEL.get(AGG_ALIASES.get(func.lower(), ''), func)})" if col else "Conteggio")))
    else:
        raise TableError("«aggregate» deve essere un elenco o una mappa.")
    for func, col, _ in out:
        if AGG_ALIASES.get(func.lower()) is None:
            raise TableError(f"Funzione di aggregazione «{func}» sconosciuta.")
        if col:
            table.index(col)
    return out


def as_list(v: Any) -> list[str]:
    if v is None or v == "":
        return []
    if isinstance(v, (list, tuple)):
        return [str(x).strip() for x in v if str(x).strip()]
    return [x.strip() for x in str(v).split(",") if x.strip()]


def group_by(table: Table, keys: list[str], aggs: list[tuple[str, str | None, str]]) -> Table:
    kidx = [table.index(k) for k in keys]
    groups: OrderedDict = OrderedDict()
    for r in table.rows:
        key = tuple(r[i] if i < len(r) else None for i in kidx)
        groups.setdefault(key, []).append(r)
    if not aggs:
        aggs = [("count", None, "Conteggio")]
    cols = [table.columns[i] for i in kidx] + [a[2] for a in aggs]
    rows = []
    for key, members in groups.items():
        vals = list(key)
        for func, col, _ in aggs:
            if col:
                ci = table.index(col)
                vals.append(aggregate(func, [m[ci] if ci < len(m) else None for m in members]))
            else:
                vals.append(len(members) if AGG_ALIASES[func.lower()] == "count" else None)
        rows.append(vals)
    return Table(cols, rows, table.truncated)


def pivot(table: Table, spec: dict) -> Table:
    if not isinstance(spec, dict):
        raise TableError("«pivot» richiede rows, columns, values (e facoltativo agg).")
    rkeys = as_list(spec.get("rows") or spec.get("righe"))
    ckey = str(spec.get("columns") or spec.get("colonne") or "").strip()
    vcol = str(spec.get("values") or spec.get("valori") or "").strip()
    func = str(spec.get("agg") or spec.get("aggregate") or ("sum" if vcol else "count"))
    total = spec.get("total", spec.get("totale", True))
    if not rkeys or not ckey:
        raise TableError("«pivot» richiede almeno rows e columns.")
    ridx = [table.index(k) for k in rkeys]
    cidx = table.index(ckey)
    vidx = table.index(vcol) if vcol else None
    col_vals: list = []
    cells: OrderedDict = OrderedDict()
    for r in table.rows:
        rk = tuple(r[i] if i < len(r) else None for i in ridx)
        ck = r[cidx] if cidx < len(r) else None
        if ck not in col_vals:
            col_vals.append(ck)
        cells.setdefault(rk, {}).setdefault(ck, []).append(r[vidx] if vidx is not None and vidx < len(r) else 1)
    try:
        col_vals.sort(key=lambda x: (x is None, x))
    except TypeError:
        col_vals.sort(key=lambda x: (x is None, str(x)))
    cols = [table.columns[i] for i in ridx] + ["(vuoto)" if c in (None, "") else _label(c) for c in col_vals]
    if total:
        cols.append("Totale")
    rows = []
    for rk, by_c in cells.items():
        vals = list(rk)
        all_vals = []
        for c in col_vals:
            vs = by_c.get(c, [])
            all_vals += vs
            vals.append(aggregate(func, vs) if vs else None)
        if total:
            vals.append(aggregate(func, all_vals))
        rows.append(vals)
    return Table(cols, rows, table.truncated)


def _label(v: Any) -> str:
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        return fmt_number(v)
    return str(v)


def sort_rows(table: Table, spec: Any) -> None:
    items = as_list(spec)
    for item in reversed(items):  # ordinamento stabile: dall'ultima chiave alla prima
        desc = item.startswith("-")
        name = item.lstrip("-+").strip()
        m = re.match(r"^(.*?)\s+(desc|asc|decrescente|crescente)$", name, re.I)
        if m:
            name, desc = m.group(1), m.group(2).lower() in ("desc", "decrescente")
        i = table.index(name)

        def key(r, i=i):
            v = r[i] if i < len(r) else None
            if v is None or v == "":
                return (2, 0)
            if isinstance(v, (int, float)):
                return (0, v)
            if isinstance(v, dt.datetime):
                return (0, v.timestamp())
            if isinstance(v, dt.date):
                return (0, dt.datetime(v.year, v.month, v.day).timestamp())
            return (1, str(v).lower())

        nonempty = [r for r in table.rows if key(r)[0] != 2]
        empty = [r for r in table.rows if key(r)[0] == 2]
        nonempty.sort(key=key, reverse=desc)
        table.rows = nonempty + empty


def select_columns(table: Table, spec: Any) -> Table:
    """columns: [A, B as Bis] oppure {A: Nuovo nome}"""
    pairs: list[tuple[str, str]] = []
    if isinstance(spec, dict):
        pairs = [(str(k), str(v) if v not in (None, "") else str(k)) for k, v in spec.items()]
    else:
        for item in as_list(spec):
            m = re.match(r"^(.*?)\s+(?:as|come)\s+(.+)$", item, re.I)
            pairs.append((m.group(1).strip(), m.group(2).strip()) if m else (item, item))
    idx = [table.index(a) for a, _ in pairs]
    return Table([b for _, b in pairs], [[r[i] if i < len(r) else None for i in idx] for r in table.rows],
                 table.truncated)


# --------------------------------------------------------------------------- formati
def fmt_number(v: float, decimals: int | None = None) -> str:
    if decimals is None:
        if isinstance(v, int) or float(v).is_integer():
            decimals = 0
        else:
            decimals = 2
    s = f"{v:,.{decimals}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_value(v: Any, fmt: str = "") -> str:
    f = (fmt or "").strip().lower()
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "Sì" if v else "No"
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y" if f == "date" else "%d/%m/%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%m/%Y" if f in ("month", "mese") else "%d/%m/%Y")
    if isinstance(v, (int, float)):
        if f in ("euro", "eur", "€"):
            return f"€ {fmt_number(v, 2)}"
        if f in ("usd", "dollar", "$"):
            return f"$ {fmt_number(v, 2)}"
        if f in ("percent", "percentuale"):
            return f"{fmt_number(v * 100, 1 if abs(v * 100 - round(v * 100)) > 1e-9 else 0)}%"
        if f in ("pct", "%"):
            return f"{fmt_number(v)}%"
        if f in ("int", "intero", "0"):
            return fmt_number(round(v), 0)
        if re.fullmatch(r"0[.,](0+)", f):
            return fmt_number(v, len(f) - 2)
        if f in ("text", "testo", "raw"):
            return str(v)
        return fmt_number(v)
    return str(v)


def sort_value(v: Any) -> str:
    if isinstance(v, bool):
        return str(int(v))
    if isinstance(v, (int, float)):
        return repr(float(v))
    if isinstance(v, dt.datetime):
        return v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return ""
